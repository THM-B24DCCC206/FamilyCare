# -*- coding: utf-8 -*-
"""
doc_common.py: Thu vien tien ich dinh dang bao cao Word chuan Do an Tot nghiep
"""
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

COLOR_PRIMARY = RGBColor(15, 23, 42)      # Slate 900
COLOR_BLUE = RGBColor(30, 58, 138)        # Blue 900
COLOR_SECONDARY = RGBColor(71, 85, 105)   # Slate 600
COLOR_TEXT = RGBColor(30, 41, 59)         # Slate 800
COLOR_RED = RGBColor(220, 38, 38)         # Red 600

def create_base_document():
    doc = Document()
    
    # Thiet lap kho giay A4 (210mm x 297mm = 8.27in x 11.69in)
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    
    # Le tieu chuan do an hoc thuat: Trai 3cm, Phai 2cm, Tren 2cm, Duoi 2cm
    section.left_margin = Inches(1.18)    # 3.0 cm
    section.right_margin = Inches(0.79)   # 2.0 cm
    section.top_margin = Inches(0.79)     # 2.0 cm
    section.bottom_margin = Inches(0.79)  # 2.0 cm
    
    # Header & Footer
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = hp.add_run("Đồ án Tốt nghiệp / Báo cáo Kỹ thuật: Hệ thống Giám sát Sức khỏe Gia đình FamilyCare")
    hrun.font.name = "Times New Roman"
    hrun.font.size = Pt(8.5)
    hrun.font.color.rgb = RGBColor(148, 163, 184)
    hrun.italic = True
    
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    frun = fp.add_run("Học viện Công nghệ Bưu chính Viễn thông – Khoa Công nghệ Thông tin")
    frun.font.name = "Times New Roman"
    frun.font.size = Pt(8.5)
    frun.font.color.rgb = RGBColor(148, 163, 184)
    
    # Configure Normal Style
    style_normal = doc.styles['Normal']
    font_normal = style_normal.font
    font_normal.name = 'Times New Roman'
    font_normal.size = Pt(13)
    font_normal.color.rgb = COLOR_TEXT
    
    return doc

def add_para(doc, text="", bold_prefix="", italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=5, line_spacing=1.3, indent=True, left_indent=None):
    """Them doan van ban chuan can deu 2 ben, co the kem tieu de in dam phia truoc va thut dau dong"""
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    
    if indent and align == WD_ALIGN_PARAGRAPH.JUSTIFY:
        p.paragraph_format.first_line_indent = Inches(0.3)  # Thụt đầu dòng 0.75 cm
        
    if left_indent is not None:
        p.paragraph_format.left_indent = left_indent
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(13)
        r_pre.font.color.rgb = COLOR_PRIMARY
        
    if text:
        r_text = p.add_run(text)
        r_text.font.name = "Times New Roman"
        r_text.font.size = Pt(13)
        r_text.font.italic = italic
        r_text.font.color.rgb = COLOR_TEXT
        
    return p

def add_h1(doc, text, page_break=True):
    """Tieu de chuong bac 1: In dam to, mau xanh dam, luon ngat sang trang moi"""
    if page_break and len(doc.paragraphs) > 1:
        doc.add_page_break()
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.keep_with_next = True
    
    run = p.add_run(text.upper())
    run.font.name = "Times New Roman"
    run.font.size = Pt(16.5)
    run.font.bold = True
    run.font.color.rgb = COLOR_BLUE
    return p

def add_h2(doc, text, left_indent=None):
    """Tieu de muc bac 2: In dam, mau xanh navy"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    if left_indent is not None:
        p.paragraph_format.left_indent = left_indent
    
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(14.5)
    run.font.bold = True
    run.font.color.rgb = COLOR_PRIMARY
    return p

def add_h3(doc, text, left_indent=Inches(0.2)):
    """Tieu de tieu muc bac 3: In dam, thut le trai 0.2 inch phan cap ro rang"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    if left_indent is not None:
        p.paragraph_format.left_indent = left_indent
    
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(13.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(30, 64, 175)
    return p

def add_h4(doc, text, left_indent=Inches(0.35)):
    """Tieu de tieu muc bac 4: In nghieng dam, thut le trai 0.35 inch"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    if left_indent is not None:
        p.paragraph_format.left_indent = left_indent
        
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.italic = True
    run.font.color.rgb = RGBColor(71, 85, 105)
    return p

def add_bullet(doc, text, bold_prefix="", left_indent=None):
    """Them dong danh sach dau cham (Bullet point)"""
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.25
    if left_indent is not None:
        p.paragraph_format.left_indent = left_indent
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(13)
        r_pre.font.color.rgb = COLOR_PRIMARY
        
    r_text = p.add_run(text)
    r_text.font.name = "Times New Roman"
    r_text.font.size = Pt(13)
    r_text.font.color.rgb = COLOR_TEXT
    return p
    r_text.font.color.rgb = COLOR_TEXT
    return p

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def configure_table_layout(table, col_widths=None):
    """Dam bao bang luon tuan thu layout fixed, khong tran le, khong rach dong giua trang va toi da 6.25 inches"""
    tblPr = table._tbl.tblPr
    tblPr.append(parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>'))
    if col_widths:
        widths_float = []
        for w in col_widths:
            if hasattr(w, 'inches'):
                widths_float.append(w.inches)
            elif hasattr(w, 'twips'):
                widths_float.append(w / 914400.0)
            else:
                widths_float.append(float(w))
        total_w = sum(widths_float)
        if total_w > 6.25:
            scale = 6.25 / total_w
            widths_float = [w * scale for w in widths_float]
            
        final_widths = [Inches(w) for w in widths_float]
        for row in table.rows:
            trPr = row._tr.get_or_add_trPr()
            trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
            for c_idx, w in enumerate(final_widths):
                if c_idx < len(row.cells):
                    row.cells[c_idx].width = w
    else:
        for row in table.rows:
            trPr = row._tr.get_or_add_trPr()
            trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def create_styled_table(doc, headers, data, col_widths=None, caption=None):
    """Tao bang bieu co style chuan y te: header xanh navy, dong xen ke, vien mong"""
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(10)
        p_cap.paragraph_format.space_after = Pt(4)
        p_cap.paragraph_format.keep_with_next = True
        r_cap = p_cap.add_run(caption)
        r_cap.bold = True
        r_cap.font.name = "Times New Roman"
        r_cap.font.size = Pt(11.5)
        r_cap.font.color.rgb = COLOR_PRIMARY
        
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    # Border mong thanh lich
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="94A3B8"/>'
        f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="475569"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)
    
    # Format Header Row
    hdr_row = table.rows[0]
    hdr_trPr = hdr_row._tr.get_or_add_trPr()
    hdr_trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    
    for i, h_text in enumerate(headers):
        cell = hdr_row.cells[i]
        set_cell_background(cell, "1E3A8A")  # Xanh navy dam
        set_cell_margins(cell, top=140, bottom=140, left=140, right=140)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h_text)
        r.bold = True
        r.font.name = "Times New Roman"
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    # Format Data Rows
    for r_idx, row_data in enumerate(data):
        row = table.rows[r_idx + 1]
        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, cell_value in enumerate(row_data):
            cell = row.cells[c_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            
            # Can giua neu la cot STT hoac Trang thai
            if c_idx == 0 or (len(headers) > 4 and c_idx == len(headers) - 1):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
            r = p.add_run(str(cell_value))
            r.font.name = "Times New Roman"
            r.font.size = Pt(10.5)
            r.font.color.rgb = COLOR_TEXT
            
    # Set layout and column widths with safety auto-scaling to max 6.25 inches
    configure_table_layout(table, col_widths)
                    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return table
                    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return table

def add_code_block(doc, code_str, caption=None):
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.paragraph_format.space_before = Pt(8)
        p_cap.paragraph_format.space_after = Pt(2)
        p_cap.paragraph_format.keep_with_next = True
        r_cap = p_cap.add_run(f"Đoạn mã: {caption}")
        r_cap.font.name = "Times New Roman"
        r_cap.font.size = Pt(11)
        r_cap.font.bold = True
        r_cap.font.color.rgb = COLOR_PRIMARY
        
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    configure_table_layout(table, [Inches(6.25)])
    cell = table.cell(0, 0)
    cell.width = Inches(6.25)
    
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="2563EB"/>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(code_str.strip())
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(30, 41, 59)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_callout(doc, text, title="GHI CHÚ HỌC THUẬT QUAN TRỌNG", box_type="info"):
    border_color = "1E3A8A" if box_type == "info" else ("DC2626" if box_type == "warning" else "16A34A")
    bg_color = "F0F4F8" if box_type == "info" else ("FEF2F2" if box_type == "warning" else "F0FDF4")
    
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    configure_table_layout(table, [Inches(6.25)])
    cell = table.cell(0, 0)
    cell.width = Inches(6.25)
    
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=140)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="32" w:space="0" w:color="{border_color}"/>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.2
    run_t = p.add_run(f"📌 {title}: ")
    run_t.bold = True
    run_t.font.name = "Times New Roman"
    run_t.font.size = Pt(11)
    run_t.font.color.rgb = RGBColor.from_string(border_color)
    
    run_b = p.add_run(text)
    run_b.font.name = "Times New Roman"
    run_b.font.size = Pt(11)
    run_b.font.italic = True
    run_b.font.color.rgb = RGBColor(51, 65, 85)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_image(doc, img_path, caption, width_inches=2.8, align=WD_ALIGN_PARAGRAPH.CENTER):
    """Chen 1 anh vao tai lieu, can giua, co chu thich Hinh ben duoi"""
    import os
    if not os.path.exists(img_path):
        print(f"Canh bao: Khong tim thay anh {img_path}")
        return
        
    p_img = doc.add_paragraph()
    p_img.alignment = align
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(3)
    p_img.paragraph_format.keep_with_next = True
    run = p_img.add_run()
    run.add_picture(img_path, width=Inches(width_inches))
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = align
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(8)
    r_cap = p_cap.add_run(caption)
    r_cap.font.name = "Times New Roman"
    r_cap.font.size = Pt(10)
    r_cap.font.italic = True
    r_cap.font.color.rgb = RGBColor(71, 85, 105)

def add_image_pair(doc, img_path1, caption1, img_path2, caption2, width_inches=2.7, overall_caption=None):
    """Chen 2 anh canh nhau (side-by-side) trong 1 bang 2 cot khong vien de tiet kiem khong gian va truc quan"""
    import os
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    configure_table_layout(table, [Inches(3.1), Inches(3.1)])
    set_table_borders(table, color="FFFFFF", sz="0")
    
    # Cell 1
    c1 = table.cell(0, 0)
    set_cell_margins(c1, top=40, bottom=40, left=40, right=40)
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(0)
    p1.paragraph_format.space_after = Pt(2)
    p1.paragraph_format.keep_with_next = True
    if os.path.exists(img_path1):
        p1.add_run().add_picture(img_path1, width=Inches(width_inches))
    p1_cap = c1.add_paragraph()
    p1_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1_cap.paragraph_format.space_before = Pt(2)
    p1_cap.paragraph_format.space_after = Pt(2)
    r1 = p1_cap.add_run(caption1)
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(9.5)
    r1.font.italic = True
    r1.font.color.rgb = RGBColor(71, 85, 105)
    
    # Cell 2
    c2 = table.cell(0, 1)
    set_cell_margins(c2, top=40, bottom=40, left=40, right=40)
    p2 = c2.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(2)
    p2.paragraph_format.keep_with_next = True
    if os.path.exists(img_path2):
        p2.add_run().add_picture(img_path2, width=Inches(width_inches))
    p2_cap = c2.add_paragraph()
    p2_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2_cap.paragraph_format.space_before = Pt(2)
    p2_cap.paragraph_format.space_after = Pt(2)
    r2 = p2_cap.add_run(caption2)
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(9.5)
    r2.font.italic = True
    r2.font.color.rgb = RGBColor(71, 85, 105)
    
    if overall_caption:
        p_ov = doc.add_paragraph()
        p_ov.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_ov.paragraph_format.space_before = Pt(4)
        p_ov.paragraph_format.space_after = Pt(8)
        r_ov = p_ov.add_run(overall_caption)
        r_ov.font.name = "Times New Roman"
        r_ov.font.size = Pt(10)
        r_ov.font.bold = True
        r_ov.font.color.rgb = COLOR_PRIMARY
    else:
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_image_placeholder(doc, caption, screen_desc="Chèn ảnh chụp màn hình giao diện ứng dụng tại đây", height_pt=90, width_inches=5.8):
    """Tao khung khoang trong tieu chuan co vien net dut de nguoi dung tu chen anh chup man hinh vao Word"""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    configure_table_layout(table, [Inches(width_inches)])
    
    cell = table.cell(0, 0)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
    
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="dashed" w:sz="8" w:space="0" w:color="94A3B8"/>'
        f'<w:left w:val="dashed" w:sz="8" w:space="0" w:color="94A3B8"/>'
        f'<w:bottom w:val="dashed" w:sz="8" w:space="0" w:color="94A3B8"/>'
        f'<w:right w:val="dashed" w:sz="8" w:space="0" w:color="94A3B8"/>'
        f'</w:tcBorders>'
    )
    cell._tc.get_or_add_tcPr().append(borders)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(height_pt // 3)
    p.paragraph_format.space_after = Pt(2)
    
    r1 = p.add_run("🖼️ [ KHOẢNG TRỐNG DÀNH ĐỂ CHÈN ẢNH CHỤP MÀN HÌNH ]\n")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(11)
    r1.font.bold = True
    r1.font.color.rgb = COLOR_BLUE
    
    r2 = p.add_run(f"» {screen_desc} «\n")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(10)
    r2.font.color.rgb = COLOR_TEXT
    
    r3 = p.add_run("(Thao tác: Nhấp chuột vào khung này -> Chọn thẻ Insert / Chèn -> Pictures / Hình ảnh -> Chọn file ảnh màn hình tương ứng)")
    r3.font.name = "Times New Roman"
    r3.font.size = Pt(9)
    r3.font.italic = True
    r3.font.color.rgb = COLOR_SECONDARY
    p.paragraph_format.space_after = Pt(height_pt // 3)
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(4)
    p_cap.paragraph_format.space_after = Pt(10)
    r_cap = p_cap.add_run(caption)
    r_cap.font.name = "Times New Roman"
    r_cap.font.size = Pt(10)
    r_cap.font.bold = True
    r_cap.font.italic = True
    r_cap.font.color.rgb = COLOR_PRIMARY
